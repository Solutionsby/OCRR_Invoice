import io
import os
from email.message import EmailMessage

from openpyxl import Workbook
from openpyxl.styles import Font

from config import config_manager as cfg
from utils import database_manager as db
from core.paths import find_month_pdfs
from file_renamer import sanitize_invoice_number
import mail_sender

MONTH_NAMES_PL = [
    "styczeń", "luty", "marzec", "kwiecień", "maj", "czerwiec",
    "lipiec", "sierpień", "wrzesień", "październik", "listopad", "grudzień",
]


def _rows_with_matched_pdfs(year: int, month: int, pdfs: list) -> list:
    """
    FAKTURY_KOSZTOWE (widok Faktury_Kosztowe_Zmapowane) za dany miesiąc
    obejmuje dziś dużo więcej niż to narzędzie kiedykolwiek widziało jako
    PDF — KSeF trafia tam też zewnętrznym kanałem niezależnym od plików na
    dysku. Zwraca pary (wiersz, ścieżka_pliku) tylko dla wierszy, których
    numer faktury (posanityzowany tak samo jak przy zmianie nazwy pliku,
    patrz file_renamer.py) pasuje do nazwy jednego z podanych PDF-ów — to
    jedyne miejsce z tą logiką dopasowania, używane zarówno do budowy
    Excela (tylko pozycje z załącznikiem) jak i do wycięcia właściwego
    podzbioru PDF-ów dla raportu wg działu/kontrahenta.
    """
    prefix = f"{month:02d}_{year}_FV_"
    pairs = []
    for r in db.get_kosztowe_by_month(year, month):
        needle = f"{prefix}{sanitize_invoice_number(r['invoice_number'])}_"
        match = next((p for p in pdfs if p.name.startswith(needle)), None)
        if match:
            pairs.append((r, match))
    return pairs


def _rows_with_attachment(year: int, month: int, pdfs: list) -> list:
    """Sama lista wierszy z _rows_with_matched_pdfs — dla build_preview,
    które nie potrzebuje wiedzieć KTÓRY plik pasuje do którego wiersza."""
    return [r for r, _ in _rows_with_matched_pdfs(year, month, pdfs)]


def _months_in_range(year_from: int, month_from: int, year_to: int, month_to: int) -> list:
    """Lista (rok, miesiąc) od "od" do "do" włącznie — raporty wg działu/
    kontrahenta mogą obejmować dowolne okno, nie tylko jeden kalendarzowy
    miesiąc (np. cały kwartał albo przełom roku). Tolerancyjne na odwrócony
    zakres (od > do) — zamieniamy końce zamiast błądzić."""
    start = year_from * 12 + (month_from - 1)
    end = year_to * 12 + (month_to - 1)
    if end < start:
        start, end = end, start
    return [(m // 12, m % 12 + 1) for m in range(start, end + 1)]


def _range_label(year_from: int, month_from: int, year_to: int, month_to: int) -> str:
    start = f"{MONTH_NAMES_PL[month_from - 1]} {year_from}"
    if (year_from, month_from) == (year_to, month_to):
        return start
    end = f"{MONTH_NAMES_PL[month_to - 1]} {year_to}"
    return f"{start} – {end}"


def _range_filename_part(year_from: int, month_from: int, year_to: int, month_to: int) -> str:
    if (year_from, month_from) == (year_to, month_to):
        return f"{month_from:02d}_{year_from}"
    return f"{month_from:02d}_{year_from}-{month_to:02d}_{year_to}"


def _scoped_rows_and_pdfs_range(year_from: int, month_from: int, year_to: int, month_to: int,
                                 predicate, with_pdfs: bool = True) -> tuple:
    """
    Wiersze (i, gdy with_pdfs, TYLKO pasujące im pliki PDF) dla predykatu
    (dział/kontrahent), zsumowane po wszystkich miesiącach zakresu.

    with_pdfs=True: tylko wiersze mające odpowiadający PDF (spójne z
    załącznikami, tak jak pełna paczka) — myśl księgowa, "chcę dokumenty
    źródłowe razem z zestawieniem".
    with_pdfs=False: WSZYSTKIE wiersze spełniające predykat, niezależnie od
    obecności PDF (tak jak raport właściciela) i bez żadnych załączników —
    myśl zarządcza, lżejszy mail z pełnym obrazem kosztów tego
    działu/kontrahenta, bez ukrywania pozycji bez pliku.
    """
    rows, pdfs = [], []
    for year, month in _months_in_range(year_from, month_from, year_to, month_to):
        if with_pdfs:
            month_pdfs = find_month_pdfs(year, month)
            pairs = [(r, p) for r, p in _rows_with_matched_pdfs(year, month, month_pdfs) if predicate(r)]
            rows.extend(r for r, _ in pairs)
            pdfs.extend(p for _, p in pairs)
        else:
            rows.extend(r for r in db.get_kosztowe_by_month(year, month) if predicate(r))
    return rows, pdfs


def _safe_filename_part(value: str) -> str:
    """Dział/kontrahent trafiają do nazwy pliku Excela — te same znaki
    zakazane co przy zmianie nazwy PDF-ów (patrz file_renamer.py)."""
    return str(value).replace("/", "_").replace("\\", "_").replace(" ", "_").replace(":", "_")


def list_dzialy(year_from: int, month_from: int, year_to: int, month_to: int) -> list:
    """Suma działów pojawiających się w KTÓRYMKOLWIEK miesiącu zakresu —
    zasila select w UI, zanim operator w ogóle wybierze dział."""
    result = set()
    for year, month in _months_in_range(year_from, month_from, year_to, month_to):
        result.update(r["dzial"] for r in db.get_kosztowe_by_month(year, month))
    return sorted(result)


def list_kontrahenci(year_from: int, month_from: int, year_to: int, month_to: int) -> list:
    result = set()
    for year, month in _months_in_range(year_from, month_from, year_to, month_to):
        result.update(r["firm_name"] for r in db.get_kosztowe_by_month(year, month))
    return sorted(result)


def build_preview(year: int, month: int) -> dict:
    """Podgląd przed wysyłką: co wejdzie do Excela (tylko pozycje z
    FAKTURY_KOSZTOWE, dla których jest odpowiadający PDF) i ile plików PDF
    zostanie załączonych (archiwum + wciąż oczekujące na zapłatę)."""
    pdfs = find_month_pdfs(year, month)
    rows = _rows_with_attachment(year, month, pdfs)
    total_brutto = round(sum(r["netto"] + r["vat"] for r in rows), 2)
    return {
        "year": year,
        "month": month,
        "invoice_count": len(rows),
        "total_brutto": total_brutto,
        "attachment_count": len(pdfs),
        "attachments": [p.name for p in pdfs],
    }


def _build_excel(rows: list, include_dzial: bool = False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    headers = ["Firma", "Nr faktury", "Data wystawienia", "Netto", "VAT", "Brutto"]
    if include_dzial:
        headers.append("Dział")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_netto = total_vat = total_brutto = 0.0
    for r in rows:
        brutto = round(r["netto"] + r["vat"], 2)
        row_values = [r["firm_name"], r["invoice_number"], r["invoice_date"], r["netto"], r["vat"], brutto]
        if include_dzial:
            row_values.append(r.get("dzial", ""))
        ws.append(row_values)
        total_netto += r["netto"]
        total_vat += r["vat"]
        total_brutto += brutto

    total_row = ["", "", "SUMA:", round(total_netto, 2), round(total_vat, 2), round(total_brutto, 2)]
    if include_dzial:
        total_row.append("")
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    widths = {"A": 30, "B": 20, "C": 16, "D": 14, "E": 14, "F": 14}
    if include_dzial:
        widths["G"] = 20
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dzial_breakdown(rows: list) -> list:
    """Suma brutto i liczba pozycji na Dział, posortowane malejąco po
    kwocie — źródło podsumowania w treści raportu właściciela."""
    totals = {}
    for r in rows:
        dzial = r.get("dzial") or "(brak działu)"
        entry = totals.setdefault(dzial, {"dzial": dzial, "count": 0, "brutto": 0.0})
        entry["count"] += 1
        entry["brutto"] += round(r["netto"] + r["vat"], 2)
    breakdown = sorted(totals.values(), key=lambda e: e["brutto"], reverse=True)
    for e in breakdown:
        e["brutto"] = round(e["brutto"], 2)
    return breakdown


def _owner_recipients() -> list:
    return cfg.load_settings().get("email_config", {}).get("owner_recipients", [])


def _send_excel_pdf_report(*, rows: list, pdfs: list, subject: str, filename_prefix: str,
                            recipients: list, intro_html: str, no_recipients_error: str) -> dict:
    """
    Wspólna budowa+wysyłka dla wariantów "Excel + PDF-y" (pełna paczka,
    raport wg działu, raport wg kontrahenta) — różnią się tylko tym, jakie
    wiersze/pliki dostają na wejściu, tematem maila i skąd biorą odbiorców.
    """
    result = {
        "sent": False, "count": 0, "attachment_count": 0,
        "recipients": recipients, "missing_files": [], "error": None,
    }
    if not rows and not pdfs:
        result["error"] = "Brak faktur dla wybranego zakresu."
        return result
    if not recipients:
        result["error"] = no_recipients_error
        return result

    conf = cfg.load_settings().get("email_config", {})
    recipient_str = ", ".join(recipients)

    sender_email = os.getenv('SENDER')
    msg = EmailMessage()
    msg['Subject'] = subject
    msg['From'] = f"Faktury — Raport <{sender_email}>"
    msg['To'] = recipient_str

    total_brutto = round(sum(r["netto"] + r["vat"] for r in rows), 2)
    attachments_line = (
        f"W załącznikach: arkusz Excel z podsumowaniem oraz {len(pdfs)} plik(ów) PDF."
        if pdfs else "W załączniku: arkusz Excel z podsumowaniem (bez PDF-ów)."
    )
    html_body = f"""
    <html>
    <body style="font-family: Arial, sans-serif;">
        <h2>{subject}</h2>
        {intro_html}
        <p>Pozycji w podsumowaniu: <strong>{len(rows)}</strong></p>
        <p>Suma brutto: <strong>{total_brutto:.2f} zł</strong></p>
        <p>{attachments_line}</p>
    </body>
    </html>
    """
    msg.add_alternative(html_body, subtype='html')

    excel_bytes = _build_excel(rows)
    msg.add_attachment(
        excel_bytes,
        maintype='application',
        subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename=f"{filename_prefix}.xlsx",
    )

    for pdf_path in pdfs:
        try:
            with open(pdf_path, 'rb') as f:
                msg.add_attachment(f.read(), maintype='application', subtype='pdf', filename=pdf_path.name)
        except Exception as e:
            print(f"⚠️ Nie udało się dołączyć {pdf_path.name}: {e}")
            result["missing_files"].append(str(pdf_path))

    try:
        smtp_port = int(conf.get('smtp_port', 587))
        print(f"⏳ Wysyłanie „{subject}” do {recipient_str}...")
        mail_sender.send_via_smtp(msg, smtp_port)
        print("✅ SUKCES: wysłano!")

        result["sent"] = True
        result["count"] = len(rows)
        result["attachment_count"] = len(pdfs) - len(result["missing_files"])
    except Exception as e:
        print(f"❌ BŁĄD: {e}")
        result["error"] = str(e)

    return result


def send_monthly_package(year: int, month: int) -> dict:
    """
    Comiesięczna pełna paczka: jeden mail z arkuszem Excel (tylko pozycje z
    FAKTURY_KOSZTOWE, dla których jest odpowiadający PDF — patrz
    _rows_with_matched_pdfs) i wszystkimi PDF-ami faktur danego miesiąca
    (wszystkie rodzaje — KSeF/inne/EUR — zarówno już zarchiwizowane, jak i
    wciąż oczekujące na zapłatę). W odróżnieniu od send_by_ids/
    send_payment_report (core/mailer.py, mail_sender.py) to wysyłka
    wyłącznie informacyjna — nie przenosi plików ani nie zmienia statusów w
    bazie, żeby nie kolidować ze stanem, którym zarządza cotygodniowy mailer
    przypomnień.
    """
    pdfs = find_month_pdfs(year, month)
    rows = _rows_with_attachment(year, month, pdfs)
    month_label = f"{MONTH_NAMES_PL[month - 1]} {year}"
    return _send_excel_pdf_report(
        rows=rows, pdfs=pdfs,
        subject=f"Pełna paczka faktur — {month_label}",
        filename_prefix=f"Podsumowanie_{month:02d}_{year}",
        recipients=mail_sender.get_recipients(),
        intro_html="",
        no_recipients_error="Brak zdefiniowanych odbiorców (settings.json: email_config.recipients).",
    )


def build_dzial_preview(year_from: int, month_from: int, year_to: int, month_to: int,
                         dzial: str, with_pdfs: bool = True) -> dict:
    """Podgląd raportu dla jednego działu w wybranym oknie miesięcy/lat."""
    rows, pdfs = _scoped_rows_and_pdfs_range(
        year_from, month_from, year_to, month_to, lambda r: r["dzial"] == dzial, with_pdfs,
    )
    total_brutto = round(sum(r["netto"] + r["vat"] for r in rows), 2)
    return {
        "dzial": dzial, "invoice_count": len(rows), "total_brutto": total_brutto,
        "attachment_count": len(pdfs), "attachments": [p.name for p in pdfs],
    }


def send_dzial_report(year_from: int, month_from: int, year_to: int, month_to: int,
                       dzial: str, recipients: list, with_pdfs: bool = True) -> dict:
    """Excel (+ PDF-y, jeśli with_pdfs) jednego działu za wybrane okno.
    Odbiorcy podawani ręcznie przy wysyłce (bez zapisanej mapy
    dział→odbiorcy)."""
    rows, pdfs = _scoped_rows_and_pdfs_range(
        year_from, month_from, year_to, month_to, lambda r: r["dzial"] == dzial, with_pdfs,
    )
    label = _range_label(year_from, month_from, year_to, month_to)
    return _send_excel_pdf_report(
        rows=rows, pdfs=pdfs,
        subject=f"Raport działu „{dzial}” — {label}",
        filename_prefix=f"Raport_{_safe_filename_part(dzial)}_{_range_filename_part(year_from, month_from, year_to, month_to)}",
        recipients=recipients,
        intro_html=f"<p>Dział: <strong>{dzial}</strong></p>",
        no_recipients_error="Podaj przynajmniej jednego odbiorcę.",
    )


def build_kontrahent_preview(year_from: int, month_from: int, year_to: int, month_to: int,
                              kontrahent: str, with_pdfs: bool = True) -> dict:
    """Podgląd raportu dla jednego kontrahenta w wybranym oknie miesięcy/lat."""
    rows, pdfs = _scoped_rows_and_pdfs_range(
        year_from, month_from, year_to, month_to, lambda r: r["firm_name"] == kontrahent, with_pdfs,
    )
    total_brutto = round(sum(r["netto"] + r["vat"] for r in rows), 2)
    return {
        "kontrahent": kontrahent, "invoice_count": len(rows), "total_brutto": total_brutto,
        "attachment_count": len(pdfs), "attachments": [p.name for p in pdfs],
    }


def send_kontrahent_report(year_from: int, month_from: int, year_to: int, month_to: int,
                            kontrahent: str, recipients: list, with_pdfs: bool = True) -> dict:
    """Excel (+ PDF-y, jeśli with_pdfs) jednego kontrahenta za wybrane okno.
    Odbiorcy podawani ręcznie przy wysyłce (bez zapisanej mapy
    kontrahent→odbiorcy)."""
    rows, pdfs = _scoped_rows_and_pdfs_range(
        year_from, month_from, year_to, month_to, lambda r: r["firm_name"] == kontrahent, with_pdfs,
    )
    label = _range_label(year_from, month_from, year_to, month_to)
    return _send_excel_pdf_report(
        rows=rows, pdfs=pdfs,
        subject=f"Raport kontrahenta „{kontrahent}” — {label}",
        filename_prefix=f"Raport_{_safe_filename_part(kontrahent)}_{_range_filename_part(year_from, month_from, year_to, month_to)}",
        recipients=recipients,
        intro_html=f"<p>Kontrahent: <strong>{kontrahent}</strong></p>",
        no_recipients_error="Podaj przynajmniej jednego odbiorcę.",
    )


def build_owner_preview(year: int, month: int) -> dict:
    """
    Podgląd raportu właściciela: CAŁA FAKTURY_KOSZTOWE za miesiąc, bez
    filtrowania po obecności PDF (w odróżnieniu od build_preview) — to
    raport zarządczy "ile wydaliśmy i na co", nie komplet dokumentów
    księgowych, więc ma pokazywać pełny obraz kosztów firmy.
    """
    rows = db.get_kosztowe_by_month(year, month)
    total_brutto = round(sum(r["netto"] + r["vat"] for r in rows), 2)
    return {
        "year": year,
        "month": month,
        "invoice_count": len(rows),
        "total_brutto": total_brutto,
        "breakdown": _dzial_breakdown(rows),
    }


def send_owner_report(year: int, month: int) -> dict:
    """
    Raport właściciela: sam Excel (cała FAKTURY_KOSZTOWE za miesiąc, z
    kolumną Dział) + krótkie podsumowanie z podziałem na Dział w treści
    maila. Celowo bez załączników PDF — to raport zarządczy, nie komplet
    dokumentów księgowych (patrz send_monthly_package dla tamtego).
    Odbiorcy: email_config.owner_recipients — niezależna lista od paczki
    dla księgowości/przypomnień.
    """
    recipients = _owner_recipients()
    result = {"sent": False, "count": 0, "recipients": recipients, "error": None}

    rows = db.get_kosztowe_by_month(year, month)
    if not rows:
        result["error"] = f"Brak faktur za {month:02d}/{year}."
        return result
    if not recipients:
        result["error"] = "Brak zdefiniowanych odbiorców raportu właściciela (settings.json: email_config.owner_recipients)."
        return result

    conf = cfg.load_settings().get("email_config", {})
    recipient_str = ", ".join(recipients)
    month_label = f"{MONTH_NAMES_PL[month - 1]} {year}"
    breakdown = _dzial_breakdown(rows)
    total_brutto = round(sum(e["brutto"] for e in breakdown), 2)

    sender_email = os.getenv('SENDER')
    msg = EmailMessage()
    msg['Subject'] = f"Raport kosztów wg działów — {month_label}"
    msg['From'] = f"Faktury — Raport Właściciela <{sender_email}>"
    msg['To'] = recipient_str

    rows_html = "".join(
        f"<tr><td style='padding:6px; border-bottom:1px solid #ddd;'>{e['dzial']}</td>"
        f"<td style='padding:6px; border-bottom:1px solid #ddd; text-align:right;'>{e['count']}</td>"
        f"<td style='padding:6px; border-bottom:1px solid #ddd; text-align:right;'>{e['brutto']:.2f} zł</td></tr>"
        for e in breakdown
    )
    html_body = f"""
    <html>
    <body style="font-family: Arial, sans-serif;">
        <h2>Raport kosztów wg działów — {month_label}</h2>
        <p>Pozycji łącznie: <strong>{len(rows)}</strong> · Suma brutto: <strong>{total_brutto:.2f} zł</strong></p>
        <table style="border-collapse: collapse; width: 100%; max-width: 480px;">
            <thead>
                <tr>
                    <th style="text-align:left; padding:6px; border-bottom:2px solid #ddd;">Dział</th>
                    <th style="text-align:right; padding:6px; border-bottom:2px solid #ddd;">Liczba faktur</th>
                    <th style="text-align:right; padding:6px; border-bottom:2px solid #ddd;">Suma brutto</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
        <p style="margin-top: 16px;">Pełne zestawienie pozycji w załączonym arkuszu Excel.</p>
    </body>
    </html>
    """
    msg.add_alternative(html_body, subtype='html')

    excel_bytes = _build_excel(rows, include_dzial=True)
    msg.add_attachment(
        excel_bytes,
        maintype='application',
        subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename=f"Raport_{month:02d}_{year}.xlsx",
    )

    try:
        smtp_port = int(conf.get('smtp_port', 587))
        print(f"⏳ Wysyłanie raportu właściciela ({month_label}) do {recipient_str}...")
        mail_sender.send_via_smtp(msg, smtp_port)
        print("✅ SUKCES: Raport właściciela wysłany!")
        result["sent"] = True
        result["count"] = len(rows)
    except Exception as e:
        print(f"❌ BŁĄD: {e}")
        result["error"] = str(e)

    return result
